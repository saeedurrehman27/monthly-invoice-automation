"""
Monthly invoice run.

Reads the client roster, groups lines per client, splits each all-in amount
into an untaxed and a taxed invoice, renders both PDFs per client, and bumps
each client's next invoice number by 2.

Usage:
    python generate_invoices.py --review                      # print the numbers, change nothing
    python generate_invoices.py --month 11 --year 2026        # generate everyone
    python generate_invoices.py --day 14                      # only clients billed on the 14th
"""
import argparse
import datetime
import json
import os

import openpyxl

from invoicing import generate_invoice, group_roster

COLUMNS = {"client": 1, "email": 2, "line_item": 3, "amount": 4,
           "item_type": 5, "invoice_day": 6, "next_invoice_no": 11}


def load_rows(ws):
    rows = []
    for r in range(2, ws.max_row + 1):
        row = {key: ws.cell(row=r, column=col).value for key, col in COLUMNS.items()}
        row["row_index"] = r
        rows.append(row)
    return rows


def print_review(groups, currency):
    """Show every number before anything is generated, so the owner can check it."""
    print(f"{'Client':<26}{'Day':>4}  {'Line':<14}{'All-in':>11}{'Untaxed':>11}{'Taxed':>10}{'Tax':>9}")
    for g in groups:
        for line in g["lines"]:
            print(f"{g['client']:<26}{g['invoice_day']:>4}  {line['desc']:<14}{line['amount']:>11,}"
                  f"{line['untaxed']:>11,}{line['taxed_subtotal']:>10,}{line['tax']:>9,}")
    print(f"\nAmounts in {currency}. Run without --review to generate the PDFs.")


def run(roster_path, config, month, year, out_dir, day_filter=None, review=False):
    company, billing = config["company"], config["billing"]
    wb = openpyxl.load_workbook(roster_path)
    ws = wb["Roster"]
    groups = group_roster(load_rows(ws), billing["untaxed_share"], billing["tax_rate"], day_filter)

    if review:
        print_review(groups, company.get("currency", ""))
        return []

    os.makedirs(out_dir, exist_ok=True)
    results = []
    for g in groups:
        date_str = f"{int(g['invoice_day']):02d}/{month:02d}/{year % 100:02d}"
        untaxed_no = f"INV-{g['next_invoice_no']:04d}"
        taxed_no = f"INV-{g['next_invoice_no'] + 1:04d}"
        safe_name = "".join(ch for ch in g["client"] if ch.isalnum() or ch in " -_").strip()

        untaxed_pdf = os.path.join(out_dir, f"{safe_name} - {billing['untaxed_invoice_label']} - {untaxed_no}.pdf")
        untaxed = generate_invoice(untaxed_pdf, company, g["client"], untaxed_no, date_str,
                                   g["untaxed_items"], tax_rate=0.0)

        taxed_pdf, taxed = None, {"total_due": 0}
        if g["taxed_items"]:
            taxed_pdf = os.path.join(out_dir, f"{safe_name} - {billing['taxed_invoice_label']} - {taxed_no}.pdf")
            # Charge the sum of each line's rounded tax, not a re-rounded tax on the
            # combined subtotal; otherwise multi-line clients can drift by a rupee.
            taxed = generate_invoice(taxed_pdf, company, g["client"], taxed_no, date_str,
                                     g["taxed_items"], tax_rate=billing["tax_rate"],
                                     tax_amount=sum(line["tax"] for line in g["lines"]))

        for r in g["row_indexes"]:
            ws.cell(row=r, column=COLUMNS["next_invoice_no"]).value = g["next_invoice_no"] + 2

        results.append({
            "client": g["client"], "email": g["email"], "date": date_str,
            "untaxed_pdf": untaxed_pdf, "taxed_pdf": taxed_pdf,
            "grand_total": untaxed["total_due"] + taxed["total_due"],
        })

    wb.save(roster_path)
    return results


if __name__ == "__main__":
    today = datetime.date.today()
    parser = argparse.ArgumentParser(description="Generate monthly client invoices from the roster.")
    parser.add_argument("--roster", default="client_roster.xlsx")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--month", type=int, default=today.month)
    parser.add_argument("--year", type=int, default=today.year)
    parser.add_argument("--out-dir", default=None, help="Defaults to invoices_<month>_<year>")
    parser.add_argument("--day", type=int, default=None, help="Only bill clients whose invoice day is N")
    parser.add_argument("--review", action="store_true", help="Print the numbers without generating anything")
    args = parser.parse_args()

    with open(args.config, encoding="utf-8") as f:
        cfg = json.load(f)
    out = args.out_dir or f"invoices_{args.month:02d}_{args.year}"

    for result in run(args.roster, cfg, args.month, args.year, out, args.day, args.review):
        print(f"{result['client']:<28} {result['date']}  total {result['grand_total']:>12,}  -> {result['untaxed_pdf']}"
              + (f", {result['taxed_pdf']}" if result["taxed_pdf"] else ""))
