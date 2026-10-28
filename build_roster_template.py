"""
Build the client roster spreadsheet.

    python build_roster_template.py                       # blank template -> client_roster.xlsx
    python build_roster_template.py --sample examples/client_roster.xlsx

Yellow cells are inputs. The grey columns are live formulas that show the
same split the generator uses, so the owner can sanity-check amounts in Excel.
"""
import argparse
import json

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = [
    "Client Name", "Email", "Line Item", "Amount (all-in, incl. tax)", "Item Type",
    "Invoice Day of Month", "Untaxed Line (calc)", "Taxed Subtotal (calc)",
    "Tax (calc)", "Taxed Total (calc)", "Next Invoice No.",
]
WIDTHS = [24, 26, 16, 14, 14, 10, 14, 14, 12, 14, 12]
INPUT_COLS = (1, 2, 3, 4, 5, 6, 11)

# Fictional clients covering every case: single office, multi-office,
# an untaxed flat fee, and a client billed mid-month.
SAMPLE_ROWS = [
    ("Bluebird Design Studio", "accounts@bluebird.example", "Office 101", 154800, "Standard", 1, 1),
    ("Northpeak Consulting", "finance@northpeak.example", "Office 203", 120000, "Standard", 1, 11),
    ("Northpeak Consulting", "finance@northpeak.example", "Office 206", 95500, "Standard", 1, 11),
    ("Northpeak Consulting", "finance@northpeak.example", "Static IP", 5000, "Untaxed Only", 1, 11),
    ("Lumen Labs", "billing@lumenlabs.example", "Office 307", 160321, "Standard", 14, 21),
    ("Riya Malik (Freelancer)", "riya@example.com", "Hot Desk", 25000, "Standard", 9, 31),
]


def build(path, sample=False, untaxed_share=0.8, tax_rate=0.16):
    divisor = untaxed_share + (1 - untaxed_share) * (1 + tax_rate)
    taxed_share = round(1 - untaxed_share, 6)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Roster"

    thin = Side(style="thin", color="B7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, header in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = PatternFill("solid", fgColor="2C4D8F")
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    ws.row_dimensions[1].height = 46

    rows = SAMPLE_ROWS if sample else []
    for r in range(2, max(22, len(rows) + 2)):
        if r - 2 < len(rows):
            client, email, item, amount, item_type, day, next_no = rows[r - 2]
            for col, value in zip(INPUT_COLS, (client, email, item, amount, item_type, day, next_no)):
                ws.cell(row=r, column=col, value=value)

        ws.cell(row=r, column=7, value=f'=IF(E{r}="Untaxed Only",D{r},IF(D{r}="","",D{r}-J{r}))')
        ws.cell(row=r, column=8, value=f'=IF(E{r}="Untaxed Only",0,IF(D{r}="","",ROUND(D{r}/{divisor:g}*{taxed_share:g},0)))')
        ws.cell(row=r, column=9, value=f'=IF(E{r}="Untaxed Only",0,IF(H{r}="","",ROUND(H{r}*{tax_rate:g},0)))')
        ws.cell(row=r, column=10, value=f'=IF(E{r}="Untaxed Only",0,IF(D{r}="","",H{r}+I{r}))')

        for col in range(1, 12):
            cell = ws.cell(row=r, column=col)
            cell.border = border
            if col in INPUT_COLS:
                cell.font = Font(name="Arial", color="0000FF")
                cell.fill = PatternFill("solid", fgColor="FFF2CC")
            else:
                cell.font = Font(name="Arial", color="000000")
                cell.fill = PatternFill("solid", fgColor="F2F2F2")
            if col in (4, 7, 8, 9, 10):
                cell.number_format = '#,##0;(#,##0);"-"'

    ws.freeze_panes = "A2"
    ws["A1"].comment = Comment(
        "One row = one billable line. Rows with the same Client Name + Email + Invoice Day "
        "become one combined untaxed invoice and one combined taxed invoice. "
        "'Standard' = Amount is the all-in total, auto-split. "
        "'Untaxed Only' = whole Amount goes on the untaxed invoice, no split, no tax.",
        "Invoice Automation",
    )
    wb.save(path)
    return path


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("path", nargs="?", default="client_roster.xlsx")
    parser.add_argument("--sample", action="store_true", help="Fill in fictional sample clients")
    parser.add_argument("--config", default=None, help="Read untaxed_share and tax_rate from a config file")
    args = parser.parse_args()

    billing = {"untaxed_share": 0.8, "tax_rate": 0.16}
    if args.config:
        with open(args.config, encoding="utf-8") as f:
            billing.update(json.load(f)["billing"])
    print("saved", build(args.path, args.sample, billing["untaxed_share"], billing["tax_rate"]))
